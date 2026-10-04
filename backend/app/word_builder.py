"""Генерация Word-документа из значений проекта."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from app.template_loader import load_manifest, load_section

EMPTY = "________________"
FONT = "Times New Roman"
SIZE = Pt(14)


# ---------- утилиты ----------

def _set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15


def _add_run(p, text: str, bold: bool = False, italic: bool = False):
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.size = SIZE
    return run


def _add_paragraph(doc, text: str = "", align=None) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        _add_run(p, text)
    return p


def _add_heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, text, bold=True)


def _val(values, key, default=None):
    v = values.get(key)
    if v is None or v == "":
        return default
    return v


def _fpsr_phrase(positions, mode="instrumental", suffix=""):
    if not positions:
        return ""
    numbers = []
    for p in positions:
        if isinstance(p, dict):
            n = str(p.get("position", "")).strip()
            if n:
                numbers.append(n)
    if not numbers:
        return ""
    joined = ", ".join(numbers)
    if mode == "instrumental":
        phrase = f"позицией {joined}" if len(numbers) == 1 else f"позициями {joined}"
    else:
        phrase = f"позиции {joined}" if len(numbers) == 1 else f"позициям {joined}"
    return phrase + suffix


# ---------- части абзаца ----------

def _render_paragraph_parts(p, parts, values, fields) -> None:
    for part in parts:
        t = part.get("type")

        if t == "text":
            _add_run(p, part.get("text", ""))

        elif t == "field":
            field_name = part.get("field")
            def_ = fields.get(field_name) or {}
            value = _val(values, field_name, def_.get("default"))
            if value in (None, ""):
                value = EMPTY if def_.get("required") else def_.get("default", "")
            _add_run(p, value if value is not None else "")

        elif t == "choice_text":
            field_name = part.get("field")
            value = _val(values, field_name, part.get("default"))
            if value is None:
                value = part.get("default", "")
            _add_run(p, value)

        elif t == "fpsr_reference":
            field_name = part.get("field")
            positions = values.get(field_name) or []
            if part.get("mode") == "readonly":
                phrase = _fpsr_phrase(positions, "dative", part.get("suffix", " ФПСР"))
                if not phrase:
                    phrase = part.get("empty_text", EMPTY)
            else:
                phrase = _fpsr_phrase(positions, "instrumental", "")
                if not phrase:
                    phrase = EMPTY
            _add_run(p, phrase)


# ---------- блоки ----------

def _render_approval_block(doc, block) -> None:
    for line_parts in block.get("lines", []):
        p = doc.add_paragraph()
        for part in line_parts:
            if part.get("type") == "text":
                _add_run(p, part.get("text", ""))


def _render_main_title(doc, title, values) -> None:
    for line_parts in title.get("lines", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for part in line_parts:
            if part.get("type") == "text":
                _add_run(p, part.get("text", ""), bold=True)
            elif part.get("type") == "field":
                _add_run(p, _val(values, part.get("field"), ""), bold=True)


def _render_heading(doc, block) -> None:
    _add_heading(doc, block.get("text", ""))


def _render_paragraph(doc, block, values, fields) -> None:
    p = doc.add_paragraph()
    _render_paragraph_parts(p, block.get("parts", []), values, fields)


def _render_field_block(doc, block, values, fields) -> None:
    field_name = block.get("field")
    def_ = fields.get(field_name) or {}
    value = _val(values, field_name, def_.get("default"))
    if value in (None, ""):
        value = EMPTY if def_.get("required") else ""
    _add_paragraph(doc, str(value or ""))


def _render_repeatable_group(doc, block, values, fields) -> None:
    field_name = block.get("field")
    def_ = fields.get(field_name) or {}
    rows = values.get(field_name) or []
    template = block.get("row_template") or {}
    parts = template.get("parts") or []
    punctuation = template.get("punctuation") or {}

    clean_rows = [
        r for r in rows
        if isinstance(r, dict) and any(str(v or "").strip() for v in r.values())
    ]
    if not clean_rows:
        return

    last_index = len(clean_rows) - 1

    for i, row in enumerate(clean_rows):
        p = doc.add_paragraph()
        for part in parts:
            pt = part.get("type")
            if pt == "text":
                _add_run(p, part.get("text", ""))
            elif pt == "field":
                _add_run(p, str(row.get(part.get("field"), "")))
            elif pt == "conditional_text":
                if part.get("condition") == "is_first_row" and i == 0:
                    _add_run(p, part.get("text", ""))

        # Знаки препинания: ';' у промежуточных, '.' у последней.
        # Если в шаблоне punctuation не задан — ничего не добавляем.
        if i < last_index and punctuation.get("middle_row"):
            _add_run(p, punctuation["middle_row"])
        elif i == last_index and punctuation.get("last_row"):
            _add_run(p, punctuation["last_row"])


def _render_classifier_paragraphs(doc, block, values, fields) -> None:
    items = values.get(block.get("field")) or []
    for item in items:
        text = item.get("paragraph") if isinstance(item, dict) else None
        if text:
            _add_paragraph(doc, text)


def _render_optional_group(doc, block, values, fields) -> None:
    controller = block.get("controller", "")
    key = controller.split(".")[-1]
    blocks = values.get("publication_blocks") or {}
    if not blocks.get(key):
        return

    field_name = block.get("field")
    def_ = fields.get(field_name) or {}
    value = _val(values, field_name, def_.get("default"))
    p = doc.add_paragraph()
    for part in block.get("parts", []):
        pt = part.get("type")
        if pt == "text":
            _add_run(p, part.get("text", ""))
        elif pt in ("field", "field_block"):
            _add_run(p, str(value or EMPTY))

def _add_formula_runs(p, formula_text: str) -> None:
    """Разбирает формулу и добавляет runs с курсивом и подстрочными/надстрочными индексами.

    Правила:
    - латинская буква -> курсив;
    - _XYZ (буквы/цифры или {}) -> подстрочный индекс (курсив);
    - ^XYZ (буквы/цифры или ()) -> надстрочный индекс;
    - цифры, знаки, пробелы, кириллица -> обычный текст.
    """
    i = 0
    n = len(formula_text)
    while i < n:
        ch = formula_text[i]

        # Подстрочный индекс
        if ch == '_':
            j = i + 1
            # _ { ... }
            if j < n and formula_text[j] == '{':
                k = formula_text.find('}', j)
                if k > 0:
                    sub = formula_text[j + 1:k]
                    r = p.add_run(sub)
                    r.font.name = FONT
                    r.font.size = SIZE
                    r.font.subscript = True
                    r.italic = True
                    i = k + 1
                    continue
            # _abc
            k = j
            while k < n and formula_text[k].isalnum():
                k += 1
            if k > j:
                sub = formula_text[j:k]
                r = p.add_run(sub)
                r.font.name = FONT
                r.font.size = SIZE
                r.font.subscript = True
                r.italic = True
                i = k
                continue
            # одиночный _
            r = p.add_run('_')
            r.font.name = FONT
            r.font.size = SIZE
            i += 1
            continue

        # Надстрочный индекс
        if ch == '^':
            j = i + 1
            if j < n and formula_text[j] == '(':
                k = formula_text.find(')', j)
                if k > 0:
                    sup = formula_text[j + 1:k]
                    r = p.add_run(sup)
                    r.font.name = FONT
                    r.font.size = SIZE
                    r.font.superscript = True
                    i = k + 1
                    continue
            k = j
            while k < n and formula_text[k].isalnum():
                k += 1
            if k > j:
                sup = formula_text[j:k]
                r = p.add_run(sup)
                r.font.name = FONT
                r.font.size = SIZE
                r.font.superscript = True
                i = k
                continue
            r = p.add_run('^')
            r.font.name = FONT
            r.font.size = SIZE
            i += 1
            continue

        # Латинская буква — курсив
        if ch.isalpha() and ch.isascii():
            r = p.add_run(ch)
            r.font.name = FONT
            r.font.size = SIZE
            r.italic = True
            i += 1
            continue

        # Всё остальное
        r = p.add_run(ch)
        r.font.name = FONT
        r.font.size = SIZE
        i += 1

def _render_formula_repeated_group(doc, block, values, fields) -> None:
    field_name = block.get("field")
    rows = values.get(field_name) or []
    indicators = values.get("indicators") or []
    clean_ind = [r for r in indicators if str(r.get("name", "")).strip()]

    point_start = block.get("point_number_start", "6.2")
    formula_num_start = block.get("formula_number_start", 1)
    try:
        major, minor = (int(x) for x in point_start.split("."))
    except Exception:
        major, minor = 6, 2

    for i, row in enumerate(rows):
        ind_name = clean_ind[i].get("name", "") if i < len(clean_ind) else ""
        okei = clean_ind[i].get("okei", "") if i < len(clean_ind) else ""
        symbol = row.get("symbol", "")
        mode = row.get("formula_mode", "site")
        formula_input = (row.get("formula_input") or "").strip()
        explanation = (row.get("explanation") or "").strip()

        point_num = f"{major}.{minor + i}."
        formula_num = formula_num_start + i

        # 1. Заголовок блока
        p = doc.add_paragraph()
        text = f"{point_num} Расчёт показателя «{ind_name}»"
        if okei:
            text += f" (код по ОКЕИ — {okei})"
        if symbol:
            text += f" ({symbol})"
        text += " осуществляется по формуле:"
        _add_run(p, text)

        # 2. Формула по центру, номер справа через табуляцию
        formula_text = (
            formula_input
            if (mode == "site" and formula_input)
            else "Требует заполнение в Word"
        )

        pf = doc.add_paragraph()
        # Настраиваем табуляцию: центр — 3.25", справа — 6.5"
        tab_stops = pf.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
        tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

        pf.add_run("\t")
        if mode == "site" and formula_input:
            _add_formula_runs(pf, formula_input)
        else:
            _add_run(pf, "Требует заполнение в Word", italic=True)
        pf.add_run("\t")
        _add_run(pf, f"({formula_num})")

        # 3. Слово «где»
        _add_paragraph(doc, "где")

                # 4. Расшифровка — с тем же разбором, что и формула
        if explanation:
            for line in explanation.split("\n"):
                line = line.strip()
                if not line:
                    continue
                pe = doc.add_paragraph()
                _add_formula_runs(pe, line)
# ---------- раздел и документ ----------

def _render_section(doc, section, values) -> None:
    fields = section.get("fields") or {}
    for block in section.get("blocks", []):
        t = block.get("type")
        if t == "heading":
            _render_heading(doc, block)
        elif t == "paragraph":
            _render_paragraph(doc, block, values, fields)
        elif t == "field_block":
            _render_field_block(doc, block, values, fields)
        elif t == "repeatable_group":
            _render_repeatable_group(doc, block, values, fields)
        elif t == "classifier_paragraphs":
            _render_classifier_paragraphs(doc, block, values, fields)
        elif t == "optional_group":
            _render_optional_group(doc, block, values, fields)
        elif t == "formula_repeated_group":
            _render_formula_repeated_group(doc, block, values, fields)
        # empty_state_output — пропускаем


def build_document(values: dict) -> BytesIO:
    doc = Document()
    _set_base_style(doc)

    manifest = load_manifest()
    sections = [load_section(s["id"]) for s in manifest["sections"]]

    opening = (sections[0].get("document_opening") or {})
    if opening.get("approval_block"):
        _render_approval_block(doc, opening["approval_block"])
        doc.add_paragraph()
    if opening.get("main_title"):
        _render_main_title(doc, opening["main_title"], values)
        doc.add_paragraph()

    for s in sections:
        _render_section(doc, s, values)
        doc.add_paragraph()

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf